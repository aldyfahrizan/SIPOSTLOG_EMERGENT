from io import BytesIO
from datetime import datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="0F172A")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14, color="0F172A")
THIN = Side(style="thin", color="CBD5E1")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _write_sheet(ws, title: str, meta: list, headers: list, rows: list):
    ws.append([title])
    ws["A1"].font = TITLE_FONT
    ws.append(["BPBD Kabupaten Banjar — SIPOSTLOG"])
    for label, value in meta:
        ws.append([f"{label}: {value}"])
    ws.append([])
    ws.append(headers)
    header_row = ws.max_row
    for col in range(1, len(headers) + 1):
        c = ws.cell(row=header_row, column=col)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
    for row in rows:
        ws.append(row)
        for col in range(1, len(headers) + 1):
            ws.cell(row=ws.max_row, column=col).border = BORDER
    widths = [len(str(h)) for h in headers]
    for row in rows:
        for i, v in enumerate(row):
            widths[i] = max(widths[i], min(len(str(v)), 48))
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w + 3
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)


def build_workbook(sheets: list) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for sheet in sheets:
        ws = wb.create_sheet(sheet["name"][:31])
        _write_sheet(ws, sheet["title"], sheet["meta"], sheet["headers"], sheet["rows"])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_opname_template(items: list, printed_by: str) -> bytes:
    headers = ["ID Item", "Nama Item", "Kategori", "Satuan", "Stok Sistem", "Stok Fisik (Isi)", "Keterangan"]
    rows = [[i["id"], i["name"], i["category"], i["unit"], i["currentStock"], None, None] for i in items]
    wb = Workbook()
    ws = wb.active
    ws.title = "Stock Opname"
    _write_sheet(
        ws,
        "Template Stock Opname — Isi kolom 'Stok Fisik (Isi)' lalu unggah kembali",
        [("Dicetak", datetime.now().strftime("%d/%m/%Y %H:%M")), ("Oleh", printed_by),
         ("Petunjuk", "Jangan ubah kolom ID Item. Kosongkan Stok Fisik jika tidak ada perubahan.")],
        headers, rows,
    )
    fill = PatternFill("solid", fgColor="FEF3C7")
    for r in range(ws.max_row - len(rows) + 1, ws.max_row + 1):
        ws.cell(row=r, column=6).fill = fill
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def parse_opname_upload(content: bytes) -> list:
    wb = load_workbook(BytesIO(content), data_only=True)
    ws = wb.active
    header_idx = None
    for row in ws.iter_rows(min_row=1, max_row=15):
        values = [str(c.value).strip().lower() if c.value is not None else "" for c in row]
        if "id item" in values:
            header_idx = row[0].row
            cols = {v: i for i, v in enumerate(values)}
            break
    if header_idx is None:
        raise ValueError("Header 'ID Item' tidak ditemukan. Gunakan template dari aplikasi.")
    id_col = cols["id item"]
    stock_col = next((i for v, i in cols.items() if v.startswith("stok fisik")), None)
    note_col = cols.get("keterangan")
    if stock_col is None:
        raise ValueError("Kolom 'Stok Fisik' tidak ditemukan.")
    entries = []
    for row in ws.iter_rows(min_row=header_idx + 1, values_only=True):
        if not row or row[id_col] in (None, ""):
            continue
        raw = row[stock_col]
        if raw in (None, ""):
            continue
        entries.append({
            "item_id": str(row[id_col]).strip(),
            "physical": raw,
            "note": (str(row[note_col]).strip() if note_col is not None and row[note_col] not in (None, "") else ""),
        })
    return entries
