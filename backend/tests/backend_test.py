"""SIPOSTLOG backend regression tests.

Covers: health, public leak-check, RBAC, transactions (IN/OUT/ADJUST),
dashboards, exports, excel template/import, and admin user management.
"""
import io
import os
import pytest
import requests
from dotenv import load_dotenv
from openpyxl import load_workbook, Workbook

load_dotenv("/app/frontend/.env", override=False)

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL").rstrip("/")
API = f"{BASE_URL}/api"

ADMIN_TOKEN = "test_session_admin_001"
PETUGAS_TOKEN = "test_session_petugas_001"
PENDING_TOKEN = "test_session_pending_001"


def h(token):
    return {"Authorization": f"Bearer {token}"}


@pytest.fixture(scope="session")
def s():
    return requests.Session()


# ---------- health & public ----------
EXPECTED_COUNT = 37


def test_health(s):
    r = s.get(f"{API}/health")
    assert r.status_code == 200
    data = r.json()
    assert data["status"] == "ok"
    assert data["items"] == EXPECTED_COUNT
    assert data["expected"] == EXPECTED_COUNT


def test_public_items_no_leak(s):
    r = s.get(f"{API}/public/items")
    assert r.status_code == 200
    items = r.json()
    assert len(items) == EXPECTED_COUNT
    banned = {
        "currentStock", "minThreshold", "current_stock", "min_threshold",
        "status", "status_counts", "stock", "quantity",
    }
    for it in items:
        assert not (banned & set(it.keys())), f"leak: {it}"
        assert set(it.keys()) == {"id", "name", "category", "unit"}


def test_public_summary_no_leak(s):
    r = s.get(f"{API}/public/summary")
    assert r.status_code == 200
    data = r.json()
    assert data["total_items"] == EXPECTED_COUNT
    assert "status_counts" not in data
    assert "status" not in data
    assert isinstance(data["categories"], list)
    assert all(set(c.keys()) == {"category", "count"} for c in data["categories"])
    # Ensure whole payload doesn't include stock fields
    assert "currentStock" not in r.text
    assert "minThreshold" not in r.text


# ---------- auth/me ----------
def test_me_no_token(s):
    r = s.get(f"{API}/auth/me")
    assert r.status_code == 401


def test_me_admin(s):
    r = s.get(f"{API}/auth/me", headers=h(ADMIN_TOKEN))
    assert r.status_code == 200
    assert r.json()["role"] == "admin"


def test_me_petugas(s):
    r = s.get(f"{API}/auth/me", headers=h(PETUGAS_TOKEN))
    assert r.status_code == 200
    assert r.json()["role"] == "petugas"


def test_me_pending(s):
    r = s.get(f"{API}/auth/me", headers=h(PENDING_TOKEN))
    assert r.status_code == 200
    assert r.json()["role"] == "pending"


# ---------- RBAC ----------
@pytest.mark.parametrize("path", ["/items", "/dashboard/stock", "/transactions"])
def test_rbac_pending_forbidden(s, path):
    r = s.get(f"{API}{path}", headers=h(PENDING_TOKEN))
    assert r.status_code == 403


def test_rbac_petugas_cannot_list_users(s):
    r = s.get(f"{API}/users", headers=h(PETUGAS_TOKEN))
    assert r.status_code == 403


def test_rbac_petugas_cannot_patch_item(s):
    # Grab any item id
    items = s.get(f"{API}/items", headers=h(ADMIN_TOKEN)).json()
    item_id = items[0]["id"]
    r = s.patch(f"{API}/items/{item_id}", headers=h(PETUGAS_TOKEN), json={"minThreshold": 5})
    assert r.status_code == 403


def test_rbac_admin_can_list_items(s):
    r = s.get(f"{API}/items", headers=h(ADMIN_TOKEN))
    assert r.status_code == 200
    assert len(r.json()) == EXPECTED_COUNT


# ---------- transactions ----------
@pytest.fixture(scope="module")
def target_item():
    r = requests.get(f"{API}/items", headers=h(ADMIN_TOKEN))
    r.raise_for_status()
    # Pick an item with enough stock for OUT test
    items = sorted(r.json(), key=lambda x: -x["currentStock"])
    return items[0]


def test_stock_in_increases(s, target_item):
    item_id = target_item["id"]
    before = s.get(f"{API}/items", headers=h(ADMIN_TOKEN)).json()
    before_qty = next(i for i in before if i["id"] == item_id)["currentStock"]
    r = s.post(f"{API}/transactions/in", headers=h(ADMIN_TOKEN), json={
        "item_id": item_id, "quantity": 10, "source": "TEST regression", "notes": "TEST_IN"
    })
    assert r.status_code == 201, r.text
    tx = r.json()
    assert tx["previous_quantity"] == before_qty
    assert tx["new_quantity"] == before_qty + 10
    assert tx["change_quantity"] == 10
    assert tx["type"] == "IN"
    after = s.get(f"{API}/items", headers=h(ADMIN_TOKEN)).json()
    after_qty = next(i for i in after if i["id"] == item_id)["currentStock"]
    assert after_qty == before_qty + 10


def test_stock_out_valid_and_over(s, target_item):
    item_id = target_item["id"]
    incident_types = s.get(f"{API}/meta/incident-types", headers=h(ADMIN_TOKEN)).json()
    inc = incident_types[0]

    # valid out
    r = s.post(f"{API}/transactions/out", headers=h(ADMIN_TOKEN), json={
        "item_id": item_id, "quantity": 3, "destination": "TEST_dest",
        "incident_type": inc, "notes": "TEST"
    })
    assert r.status_code == 201, r.text
    assert r.json()["type"] == "OUT"

    # over-quantity
    r2 = s.post(f"{API}/transactions/out", headers=h(ADMIN_TOKEN), json={
        "item_id": item_id, "quantity": 999999, "destination": "TEST_dest",
        "incident_type": inc, "notes": ""
    })
    assert r2.status_code == 400
    assert "tidak mencukupi" in r2.json()["detail"].lower()

    # bad incident type
    r3 = s.post(f"{API}/transactions/out", headers=h(ADMIN_TOKEN), json={
        "item_id": item_id, "quantity": 1, "destination": "TEST_dest",
        "incident_type": "INVALID_TYPE_XX", "notes": ""
    })
    assert r3.status_code == 400


def test_adjust_short_reason_422(s, target_item):
    r = s.post(f"{API}/transactions/adjust", headers=h(ADMIN_TOKEN), json={
        "item_id": target_item["id"], "new_quantity": 100, "reason": "ab"
    })
    assert r.status_code == 422


def test_adjust_valid(s, target_item):
    item_id = target_item["id"]
    current = next(i for i in s.get(f"{API}/items", headers=h(ADMIN_TOKEN)).json() if i["id"] == item_id)["currentStock"]
    new_qty = current + 5
    r = s.post(f"{API}/transactions/adjust", headers=h(ADMIN_TOKEN), json={
        "item_id": item_id, "new_quantity": new_qty, "reason": "TEST regression adjust"
    })
    assert r.status_code == 201, r.text
    assert r.json()["type"] == "ADJUSTMENT"
    assert r.json()["new_quantity"] == new_qty


# ---------- transactions listing & dashboards ----------
def test_list_transactions_filters(s, target_item):
    r = s.get(f"{API}/transactions", headers=h(ADMIN_TOKEN), params={"type": "IN", "limit": 10})
    assert r.status_code == 200
    for t in r.json():
        assert t["type"] == "IN"

    r2 = s.get(f"{API}/transactions", headers=h(ADMIN_TOKEN), params={"item_id": target_item["id"], "limit": 10})
    assert r2.status_code == 200
    for t in r2.json():
        assert t["item_id"] == target_item["id"]


def test_dashboard_stock(s):
    r = s.get(f"{API}/dashboard/stock", headers=h(ADMIN_TOKEN))
    assert r.status_code == 200
    d = r.json()
    assert d["total_items"] == EXPECTED_COUNT
    assert "status_counts" in d and "recent_transactions" in d and "weekly_activity" in d


def test_dashboard_distribution(s):
    r = s.get(f"{API}/dashboard/distribution", headers=h(ADMIN_TOKEN),
              params={"start": "2025-01-01", "end": "2026-12-31"})
    assert r.status_code == 200
    d = r.json()
    for k in ("per_item", "per_destination", "per_incident", "daily", "recent"):
        assert k in d


def test_dashboard_distribution_invalid_range(s):
    r = s.get(f"{API}/dashboard/distribution", headers=h(ADMIN_TOKEN),
              params={"start": "2026-05-01", "end": "2026-01-01"})
    assert r.status_code == 400


# ---------- excel exports/template ----------
@pytest.mark.parametrize("path", ["/export/stock", "/export/distribution", "/export/transactions", "/excel/template"])
def test_exports_valid_xlsx(s, path):
    r = s.get(f"{API}{path}", headers=h(ADMIN_TOKEN))
    assert r.status_code == 200, path
    assert "spreadsheet" in r.headers.get("content-type", ""), path
    wb = load_workbook(io.BytesIO(r.content))
    assert len(wb.sheetnames) >= 1


# ---------- excel import ----------
def test_excel_import_flow(s):
    # download template
    r = s.get(f"{API}/excel/template", headers=h(ADMIN_TOKEN))
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb.active

    # locate header row
    header_row = None
    header_map = {}
    for row in ws.iter_rows(min_row=1, max_row=15):
        vals = [str(c.value).strip().lower() if c.value is not None else "" for c in row]
        if "id item" in vals:
            header_row = row[0].row
            header_map = {v: i + 1 for i, v in enumerate(vals)}
            break
    assert header_row is not None
    id_col = header_map["id item"]
    sys_col = next(c for v, c in header_map.items() if v.startswith("stok sistem"))
    phys_col = next(c for v, c in header_map.items() if v.startswith("stok fisik"))

    # find first data row
    first_data = header_row + 1
    r1_id = ws.cell(row=first_data, column=id_col).value
    r1_sys = ws.cell(row=first_data, column=sys_col).value
    r2_id = ws.cell(row=first_data + 1, column=id_col).value
    r2_sys = ws.cell(row=first_data + 1, column=sys_col).value

    # row1: changed (+2), row2: unchanged (=system)
    ws.cell(row=first_data, column=phys_col, value=int(r1_sys) + 2)
    ws.cell(row=first_data + 1, column=phys_col, value=int(r2_sys))
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    up = s.post(f"{API}/excel/import", headers=h(ADMIN_TOKEN),
                files={"file": ("opname.xlsx", buf.getvalue(),
                                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert up.status_code == 200, up.text
    data = up.json()
    assert data["processed"] == 2
    assert len(data["updated"]) == 1
    assert data["unchanged"] == 1
    assert data["updated"][0]["item_id"] == r1_id


def test_excel_import_bad_extension(s):
    r = s.post(f"{API}/excel/import", headers=h(ADMIN_TOKEN),
               files={"file": ("bad.txt", b"nope", "text/plain")})
    assert r.status_code == 400


# ---------- users (admin) ----------
def test_users_list_admin(s):
    r = s.get(f"{API}/users", headers=h(ADMIN_TOKEN))
    assert r.status_code == 200
    users = r.json()
    assert any(u["user_id"] == "test-admin-001" for u in users)


def test_users_role_toggle_and_reset(s):
    # flip pending->petugas then back to pending
    r1 = s.patch(f"{API}/users/test-pending-001", headers=h(ADMIN_TOKEN), json={"role": "petugas"})
    assert r1.status_code == 200
    assert r1.json()["role"] == "petugas"
    r2 = s.patch(f"{API}/users/test-pending-001", headers=h(ADMIN_TOKEN), json={"role": "pending"})
    assert r2.status_code == 200
    assert r2.json()["role"] == "pending"
    # restore pending session (patch to pending deletes sessions)
    import pymongo
    from datetime import datetime, timedelta, timezone
    mongo_url = os.environ.get("MONGO_URL")
    db_name = os.environ.get("DB_NAME")
    assert mongo_url and db_name, "MONGO_URL and DB_NAME must be configured"
    mc = pymongo.MongoClient(mongo_url)
    db = mc[db_name]
    db.user_sessions.update_one(
        {"session_token": "test_session_pending_001"},
        {"$set": {
            "user_id": "test-pending-001",
            "session_token": "test_session_pending_001",
            "expires_at": datetime.now(timezone.utc) + timedelta(days=365),
            "created_at": datetime.now(timezone.utc),
        }},
        upsert=True,
    )


def test_users_invalid_role(s):
    r = s.patch(f"{API}/users/test-pending-001", headers=h(ADMIN_TOKEN), json={"role": "godmode"})
    assert r.status_code == 400


def test_users_cannot_demote_self(s):
    r = s.patch(f"{API}/users/test-admin-001", headers=h(ADMIN_TOKEN), json={"role": "petugas"})
    assert r.status_code == 400
    r2 = s.patch(f"{API}/users/test-admin-001", headers=h(ADMIN_TOKEN), json={"active": False})
    assert r2.status_code == 400


# ---------- item patch (admin) ----------
def test_patch_item_min_threshold(s):
    items = s.get(f"{API}/items", headers=h(ADMIN_TOKEN)).json()
    it = items[0]
    original = it["minThreshold"]
    r = s.patch(f"{API}/items/{it['id']}", headers=h(ADMIN_TOKEN), json={"minThreshold": original + 1})
    assert r.status_code == 200
    assert r.json()["minThreshold"] == original + 1
    # restore
    s.patch(f"{API}/items/{it['id']}", headers=h(ADMIN_TOKEN), json={"minThreshold": original})


# ---------- new: year param on items ----------
def test_items_year_2026_default(s):
    r = s.get(f"{API}/items", headers=h(ADMIN_TOKEN))
    assert r.status_code == 200
    items = r.json()
    assert len(items) == EXPECTED_COUNT
    # item-037 should be aman (3000 stock)
    it37 = next(i for i in items if i["id"] == "item-037")
    assert it37["status"] == "aman"
    assert it37["planQuantity"] == 3000


def test_items_year_2027_plan_view(s):
    r = s.get(f"{API}/items", headers=h(ADMIN_TOKEN), params={"year": "2027"})
    assert r.status_code == 200
    items = r.json()
    assert len(items) == EXPECTED_COUNT
    statuses = {i["status"] for i in items}
    assert statuses <= {"rencana", "tidak-dianggarkan"}
    # 4 items should be tidak-dianggarkan for 2027
    not_budgeted = [i["id"] for i in items if i["status"] == "tidak-dianggarkan"]
    assert set(not_budgeted) == {"item-015", "item-017", "item-027", "item-037"}


def test_dashboard_stock_year_2026(s):
    r = s.get(f"{API}/dashboard/stock", headers=h(ADMIN_TOKEN))
    assert r.status_code == 200
    d = r.json()
    assert d["year"] == "2026"
    assert d["realized"] is True
    assert d["total_items"] == EXPECTED_COUNT
    assert set(d["status_counts"].keys()) == {"aman", "menipis", "habis"}


def test_dashboard_stock_year_2027(s):
    r = s.get(f"{API}/dashboard/stock", headers=h(ADMIN_TOKEN), params={"year": "2027"})
    assert r.status_code == 200
    d = r.json()
    assert d["year"] == "2027"
    assert d["realized"] is False
    assert d["total_items"] == EXPECTED_COUNT
    assert set(d["status_counts"].keys()) == {"rencana", "tidak-dianggarkan"}
    assert d["status_counts"]["tidak-dianggarkan"] == 4
    assert d["status_counts"]["rencana"] == 33
    assert d["total_planned_quantity"] > 0


# ---------- new: item history chart ----------
def test_item_history_chart(s, target_item):
    r = s.get(f"{API}/items/{target_item['id']}/history-chart", headers=h(ADMIN_TOKEN))
    assert r.status_code == 200
    data = r.json()
    assert data["item_id"] == target_item["id"]
    assert "item_name" in data and "unit" in data
    assert isinstance(data["points"], list)


def test_item_history_chart_404(s):
    r = s.get(f"{API}/items/nonexistent-id/history-chart", headers=h(ADMIN_TOKEN))
    assert r.status_code == 404


# ---------- new: PDF exports ----------
@pytest.mark.parametrize("path", [
    "/export/stock/pdf",
    "/export/distribution/pdf?start=2025-01-01&end=2026-12-31",
    "/export/transactions/pdf?start=2025-01-01&end=2026-12-31",
])
def test_pdf_exports(s, path):
    r = s.get(f"{API}{path}", headers=h(ADMIN_TOKEN))
    assert r.status_code == 200, path
    assert r.headers.get("content-type", "").startswith("application/pdf"), path
    # magic bytes
    assert r.content[:4] == b"%PDF", f"{path} did not return PDF bytes"
    assert len(r.content) > 1000, f"{path} PDF suspiciously small"


def test_pdf_export_pending_forbidden(s):
    r = s.get(f"{API}/export/stock/pdf", headers=h(PENDING_TOKEN))
    assert r.status_code == 403


# ---------- new: CORS regex ----------
def test_cors_preflight_no_wildcard_with_credentials(s):
    # simulate preview subdomain
    origin = "https://logistics-hub-1573.preview.emergentagent.com"
    r = s.options(f"{API}/auth/session", headers={
        "Origin": origin,
        "Access-Control-Request-Method": "POST",
        "Access-Control-Request-Headers": "content-type",
    })
    # Must not be '*' when credentials are true
    aco = r.headers.get("access-control-allow-origin", "")
    acc = r.headers.get("access-control-allow-credentials", "")
    if acc.lower() == "true":
        assert aco != "*", f"Invalid CORS: ACAO='*' with credentials=true (origin={origin})"
