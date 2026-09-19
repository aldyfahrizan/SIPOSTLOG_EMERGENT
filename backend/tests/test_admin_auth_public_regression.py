"""Regression tests for admin login/session security, public data safety, year views, exports.

Modules/features covered in this file:
- /api/auth admin local login + session cookie security + throttle + origin guard
- /api/public payload leak prevention
- /api/dashboard stock 2026/2027 consistency
- /api/users response sanitization and self-protection
- PDF/Excel export smoke and anonymous access guard
"""

import hashlib
import io
import os
import uuid
import asyncio
import sys
from pathlib import Path
from datetime import datetime, timezone

import pytest
import pymongo
import requests
from openpyxl import load_workbook
from dotenv import load_dotenv

sys.path.append(str(Path(__file__).resolve().parents[1]))

ROOT = Path(__file__).resolve().parents[2]
load_dotenv(ROOT / "frontend" / ".env")
load_dotenv(ROOT / "backend" / ".env")

from admin_login import seed_local_admin
from motor.motor_asyncio import AsyncIOMotorClient


BASE_URL = os.environ.get("REACT_APP_BACKEND_URL")
MONGO_URL = os.environ.get("MONGO_URL")
DB_NAME = os.environ.get("DB_NAME")

if not BASE_URL:
    pytest.skip("REACT_APP_BACKEND_URL not set", allow_module_level=True)
if not MONGO_URL or not DB_NAME:
    pytest.skip("MONGO_URL or DB_NAME not set", allow_module_level=True)

API = f"{BASE_URL.rstrip('/')}/api"


@pytest.fixture(scope="session")
def http():
    return requests.Session()


@pytest.fixture(scope="session")
def mongo_db():
    client = pymongo.MongoClient(MONGO_URL)
    db = client[DB_NAME]
    yield db
    client.close()


def login_admin(http_session: requests.Session, username="admin", password="admin"):
    return http_session.post(f"{API}/auth/admin/login", json={"username": username, "password": password})


def test_admin_login_success_sets_secure_cookie_and_no_store(http):
    r = login_admin(http)
    assert r.status_code == 200, r.text
    data = r.json()
    assert data["role"] == "admin"
    assert data["active"] is True
    assert "password_hash" not in data
    assert "session_token" not in data
    assert "token" not in data
    assert "no-store" in (r.headers.get("Cache-Control") or "")

    set_cookie = r.headers.get("set-cookie", "")
    assert "session_token=" in set_cookie
    assert "HttpOnly" in set_cookie
    assert "Secure" in set_cookie
    assert "samesite=none" in set_cookie.lower()


def test_auth_me_with_cookie_and_no_store(http):
    r_login = login_admin(http)
    assert r_login.status_code == 200

    r = http.get(f"{API}/auth/me")
    assert r.status_code == 200, r.text
    data = r.json()
    assert data["role"] == "admin"
    assert "password_hash" not in data
    assert "session_token" not in data
    assert "token" not in data
    assert "no-store" in (r.headers.get("Cache-Control") or "")


def test_admin_login_wrong_credentials_and_blank_validation(http):
    wrong = login_admin(http, username="admin", password="wrong")
    assert wrong.status_code == 401
    assert "salah" in wrong.text.lower()

    blank = http.post(f"{API}/auth/admin/login", json={"username": "", "password": ""})
    assert blank.status_code == 422


def test_unauthorized_internal_endpoints_return_401():
    anon = requests.Session()
    r_items = anon.get(f"{API}/items")
    r_users = anon.get(f"{API}/users")
    assert r_items.status_code == 401
    assert r_users.status_code == 401


def test_foreign_origin_post_rejected_403():
    s = requests.Session()
    r = s.post(
        f"{API}/auth/admin/login",
        headers={"Origin": "https://evil.example"},
        json={"username": "admin", "password": "admin"},
    )
    assert r.status_code == 403


def test_allowed_origin_preflight_has_explicit_origin_and_credentials():
    origin = BASE_URL.rstrip("/")
    r = requests.options(
        f"{API}/auth/admin/login",
        headers={
            "Origin": origin,
            "Access-Control-Request-Method": "POST",
            "Access-Control-Request-Headers": "content-type",
        },
    )
    assert r.status_code in (200, 204)
    assert r.headers.get("access-control-allow-credentials", "").lower() == "true"
    assert r.headers.get("access-control-allow-origin") == origin


def test_nonexistent_user_throttle_after_five_failures(http, mongo_db):
    before_ids = {doc["_id"] for doc in mongo_db.login_attempts.find({}, {"_id": 1})}
    username = f"ghost_{uuid.uuid4().hex[:8]}"

    statuses = []
    for _ in range(5):
        rr = login_admin(http, username=username, password="invalid-pass")
        statuses.append(rr.status_code)
    sixth = login_admin(http, username=username, password="invalid-pass")
    seventh = login_admin(http, username=username, password="invalid-pass")

    assert statuses == [401, 401, 401, 401, 401]
    assert sixth.status_code in (401, 429)
    assert seventh.status_code == 429

    # Cleanup only new throttle docs created by this test
    after_docs = list(mongo_db.login_attempts.find({}, {"_id": 1}))
    new_ids = [d["_id"] for d in after_docs if d["_id"] not in before_ids]
    if new_ids:
        mongo_db.login_attempts.delete_many({"_id": {"$in": new_ids}})


def test_local_password_session_stored_as_sha256_hash(http, mongo_db):
    r = login_admin(http)
    assert r.status_code == 200
    me = http.get(f"{API}/auth/me")
    assert me.status_code == 200
    user_id = me.json()["user_id"]

    token = http.cookies.get("session_token")
    assert token and len(token) > 20

    latest = mongo_db.user_sessions.find_one({"user_id": user_id, "auth_method": "password"}, sort=[("created_at", -1)])
    assert latest is not None
    stored = latest["session_token"]
    assert stored != token
    assert len(stored) == 64
    assert stored == hashlib.sha256(token.encode()).hexdigest()


def test_logout_clears_session_and_replay_is_rejected(http):
    login = login_admin(http)
    assert login.status_code == 200
    token = http.cookies.get("session_token")
    assert token

    out = http.post(f"{API}/auth/logout")
    assert out.status_code == 200

    replay = requests.get(f"{API}/auth/me", headers={"Authorization": f"Bearer {token}"})
    assert replay.status_code == 401


def test_public_endpoints_do_not_leak_internal_stock_fields():
    items = requests.get(f"{API}/public/items")
    assert items.status_code == 200
    rows = items.json()
    assert len(rows) == 37
    for row in rows:
        assert "currentStock" not in row
        assert "minThreshold" not in row
        assert "planYear" not in row
        assert "planQuantity" not in row
        assert "transactions" not in row

    summary = requests.get(f"{API}/public/summary")
    assert summary.status_code == 200
    payload = summary.text
    assert "currentStock" not in payload
    assert "minThreshold" not in payload
    assert "planYear" not in payload
    assert "planQuantity" not in payload
    assert "transactions" not in payload


def test_dashboard_2026_2027_catalog_consistency_and_plan_rules(http):
    login = login_admin(http)
    assert login.status_code == 200

    y2026 = http.get(f"{API}/dashboard/stock", params={"year": "2026"})
    y2027 = http.get(f"{API}/dashboard/stock", params={"year": "2027"})
    assert y2026.status_code == 200
    assert y2027.status_code == 200

    d26 = y2026.json()
    d27 = y2027.json()
    assert d26["total_items"] == 37
    assert d27["total_items"] == 37
    assert "total_categories" in d27

    map26 = {i["id"]: i for i in d26["items"]}
    map27 = {i["id"]: i for i in d27["items"]}
    assert set(map26.keys()) == set(map27.keys())

    for item_id, i27 in map27.items():
        i26 = map26[item_id]
        assert i27["currentStock"] == 0
        assert i27["planQuantity"] == i26["planYear"].get("2027", 0)
        p26 = i26["planYear"].get("2026", 0)
        p27 = i26["planYear"].get("2027", 0)
        assert not (p26 > 0 and p27 > 0), f"Plan overlap in both years for {item_id}"


def test_users_and_patch_response_do_not_leak_hash_or_token(http):
    login = login_admin(http)
    assert login.status_code == 200

    users = http.get(f"{API}/users")
    assert users.status_code == 200
    rows = users.json()
    assert len(rows) >= 1
    for u in rows:
        assert "password_hash" not in u
        assert "session_token" not in u
        assert "token" not in u

    target = next(u for u in rows if u["user_id"] != login.json()["user_id"])
    original_role = target["role"]
    new_role = "petugas" if original_role != "petugas" else "pending"

    patched = http.patch(f"{API}/users/{target['user_id']}", json={"role": new_role})
    assert patched.status_code == 200, patched.text
    pdata = patched.json()
    assert pdata["role"] == new_role
    assert "password_hash" not in pdata
    assert "session_token" not in pdata
    assert "token" not in pdata

    rollback = http.patch(f"{API}/users/{target['user_id']}", json={"role": original_role})
    assert rollback.status_code == 200


def test_self_demotion_and_self_deactivation_blocked(http):
    login = login_admin(http)
    assert login.status_code == 200
    me = login.json()

    demote = http.patch(f"{API}/users/{me['user_id']}", json={"role": "petugas"})
    deactivate = http.patch(f"{API}/users/{me['user_id']}", json={"active": False})
    assert demote.status_code == 400
    assert deactivate.status_code == 400


@pytest.mark.parametrize(
    "path",
    [
        "/export/stock/pdf",
        "/export/distribution/pdf?start=2026-01-01&end=2026-12-31",
        "/export/transactions/pdf?start=2026-01-01&end=2026-12-31",
    ],
)
def test_pdf_exports_signature_and_content_smoke(http, path):
    login = login_admin(http)
    assert login.status_code == 200
    r = http.get(f"{API}{path}")
    assert r.status_code == 200, path
    assert r.headers.get("content-type", "").startswith("application/pdf")
    assert r.content[:4] == b"%PDF"
    assert len(r.content) > 1500


def test_pdf_anonymous_access_rejected():
    anon = requests.Session()
    for path in ["/export/stock/pdf", "/export/distribution/pdf", "/export/transactions/pdf"]:
        r = anon.get(f"{API}{path}")
        assert r.status_code == 401


def test_excel_template_and_import_roundtrip_non_destructive(http):
    login = login_admin(http)
    assert login.status_code == 200

    t = http.get(f"{API}/excel/template")
    assert t.status_code == 200
    assert "spreadsheet" in t.headers.get("content-type", "")

    wb = load_workbook(io.BytesIO(t.content))
    ws = wb.active

    header_row = None
    cols = {}
    for row in ws.iter_rows(min_row=1, max_row=20):
        vals = [str(c.value).strip().lower() if c.value is not None else "" for c in row]
        if "id item" in vals:
            header_row = row[0].row
            cols = {v: i + 1 for i, v in enumerate(vals)}
            break
    assert header_row is not None

    c_id = cols["id item"]
    c_sys = next(c for v, c in cols.items() if v.startswith("stok sistem"))
    c_phy = next(c for v, c in cols.items() if v.startswith("stok fisik"))

    data_row = header_row + 1
    ws.cell(row=data_row, column=c_phy, value=int(ws.cell(row=data_row, column=c_sys).value))

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    upload = http.post(
        f"{API}/excel/import",
        files={"file": ("opname.xlsx", out.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert upload.status_code == 200, upload.text
    result = upload.json()
    assert result["processed"] >= 1
    assert result["unchanged"] >= 1


def test_seed_local_admin_idempotent_no_user_or_item_reset(mongo_db):
    user_before = mongo_db.users.find_one({"username": "admin"}, {"_id": 0, "user_id": 1, "created_at": 1, "role": 1})
    item_count_before = mongo_db.items.count_documents({})
    assert user_before is not None

    async def call_seed_twice():
        client = AsyncIOMotorClient(MONGO_URL)
        try:
            adb = client[DB_NAME]
            await seed_local_admin(adb)
            await seed_local_admin(adb)
        finally:
            client.close()

    asyncio.run(call_seed_twice())

    user_after = mongo_db.users.find_one({"username": "admin"}, {"_id": 0, "user_id": 1, "created_at": 1, "role": 1})
    item_count_after = mongo_db.items.count_documents({})
    assert user_after is not None
    assert user_after["user_id"] == user_before["user_id"]
    assert user_after["created_at"] == user_before["created_at"]
    assert user_after["role"] == user_before["role"] == "admin"
    assert item_count_after == item_count_before
